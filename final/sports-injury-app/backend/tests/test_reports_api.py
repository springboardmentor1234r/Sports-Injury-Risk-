"""
tests/test_reports_api.py
----------------------------
Integration tests for the PDF and Excel export endpoints. Verifies real
file bytes come back with correct headers/content-type -- not just a 200,
since a broken PDF/Excel generator could still return 200 with garbage
bytes.
"""

import io
from openpyxl import load_workbook
from fastapi.testclient import TestClient


def register_and_login(client, email, role, password="pass1234"):
    r = client.post("/auth/register", json={"full_name": email.split("@")[0], "email": email, "password": password, "role": role})
    assert r.status_code == 201, r.text
    r = client.post("/auth/login", data={"username": email, "password": password})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


def test_pdf_export_requires_an_existing_assessment(client: TestClient):
    athlete_token = register_and_login(client, "noassess@test.com", "athlete")
    ah = {"Authorization": f"Bearer {athlete_token}"}
    profile_id = client.get("/athletes/me", headers=ah).json()["id"]

    r = client.get(f"/athletes/{profile_id}/reports/risk-assessment.pdf", headers=ah)
    assert r.status_code == 404


def test_pdf_export_returns_a_real_pdf(client: TestClient):
    athlete_token = register_and_login(client, "pdfexport@test.com", "athlete")
    physio_token = register_and_login(client, "physio_pdf@test.com", "physiotherapist")
    ah = {"Authorization": f"Bearer {athlete_token}"}
    ph = {"Authorization": f"Bearer {physio_token}"}
    profile_id = client.get("/athletes/me", headers=ah).json()["id"]

    client.post(f"/athletes/{profile_id}/risk-assessments", headers=ph)

    r = client.get(f"/athletes/{profile_id}/reports/risk-assessment.pdf", headers=ah)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:5] == b"%PDF-"  # real PDF file signature, not just "some bytes"
    assert len(r.content) > 1000


def test_excel_export_returns_a_real_workbook_with_correct_data(client: TestClient):
    athlete_token = register_and_login(client, "excelexport@test.com", "athlete")
    coach_token = register_and_login(client, "coach_excel@test.com", "coach")
    ah = {"Authorization": f"Bearer {athlete_token}"}
    ch = {"Authorization": f"Bearer {coach_token}"}
    profile_id = client.get("/athletes/me", headers=ah).json()["id"]

    client.post(f"/athletes/{profile_id}/training-load", headers=ah, json={
        "session_date": "2026-07-10", "session_type": "Conditioning", "duration_minutes": 45, "intensity_rpe": 6,
    })

    r = client.get(f"/athletes/{profile_id}/reports/training-load.xlsx", headers=ch)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert any(row[1] == "Conditioning" and row[2] == 45 for row in rows if row[1])


def test_excel_export_with_zero_entries_still_returns_valid_file(client: TestClient):
    athlete_token = register_and_login(client, "emptyexport@test.com", "athlete")
    ah = {"Authorization": f"Bearer {athlete_token}"}
    profile_id = client.get("/athletes/me", headers=ah).json()["id"]

    r = client.get(f"/athletes/{profile_id}/reports/training-load.xlsx", headers=ah)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.active is not None


def test_other_athlete_cannot_export_someone_elses_report(client: TestClient):
    athlete_token = register_and_login(client, "owner_rep@test.com", "athlete")
    ah = {"Authorization": f"Bearer {athlete_token}"}
    profile_id = client.get("/athletes/me", headers=ah).json()["id"]

    other_token = register_and_login(client, "intruder_rep@test.com", "athlete")
    oh = {"Authorization": f"Bearer {other_token}"}
    r = client.get(f"/athletes/{profile_id}/reports/training-load.xlsx", headers=oh)
    assert r.status_code == 403

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Palette
    C_BG = RGBColor(11, 17, 32)         # #0B1120 Deep Navy/Dark Slate
    C_CARD = RGBColor(30, 41, 59)       # #1E293B Card Background
    C_BORDER = RGBColor(51, 65, 85)     # #334155 Slate Border
    C_BLUE = RGBColor(56, 189, 248)     # #38BDF8 Sky/Neon Blue
    C_GREEN = RGBColor(52, 211, 153)    # #34D399 Emerald Green
    C_AMBER = RGBColor(251, 191, 36)    # #FBBF24 Amber
    C_RED = RGBColor(248, 113, 113)     # #F87171 Coral Red
    C_WHITE = RGBColor(248, 250, 252)   # #F8FAFC Heading White
    C_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Secondary Text
    C_CARD_LIGHT = RGBColor(241, 245, 249) # For light elements
    
    blank_slide_layout = prs.slide_layouts[6]
    
    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background() # No border
        return bg

    def add_header(slide, badge_text, title_text, slide_num, total_slides=13):
        # Badge
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(8), Inches(0.35))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p0 = tf.paragraphs[0]
        p0.text = badge_text.upper()
        p0.font.name = 'Arial'
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE
        
        # Title
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = 'Arial'
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_WHITE
        p1.space_before = Pt(4)
        
        # Slide number indicator
        num_box = slide.shapes.add_textbox(Inches(10.8), Inches(0.45), Inches(1.8), Inches(0.4))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.alignment = PP_ALIGN.RIGHT
        np.text = f"Slide {slide_num} of {total_slides}"
        np.font.name = 'Arial'
        np.font.size = Pt(11)
        np.font.bold = True
        np.font.color.rgb = C_BLUE
        
        # Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = C_BORDER
        line.line.fill.background()

    def add_card(slide, left, top, width, height, title, items, title_color=C_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.22)
        tf.margin_right = Inches(0.22)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)
        
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Arial'
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = title_color
        p_title.space_after = Pt(8)
        
        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.name = 'Arial'
            p.font.size = Pt(11)
            p.font.color.rgb = C_MUTED
            p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)
    
    # Title Box
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(3.0))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p_badge = tf1.paragraphs[0]
    p_badge.text = "INDIVIDUAL PROJECT PRESENTATION"
    p_badge.font.name = 'Arial'
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = C_BLUE
    p_badge.space_after = Pt(8)
    
    p_main = tf1.add_paragraph()
    p_main.text = "KineticGuard: AI-Powered Sports Injury Risk Assessment Platform"
    p_main.font.name = 'Arial'
    p_main.font.size = Pt(32)
    p_main.font.bold = True
    p_main.font.color.rgb = C_WHITE
    p_main.space_after = Pt(12)
    
    p_sub = tf1.add_paragraph()
    p_sub.text = "A Real-Time Markerless Computer Vision & Biomechanical Telemetry Pipeline for Non-Contact Injury Prevention"
    p_sub.font.name = 'Arial'
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = C_MUTED
    
    # 3 Summary Cards
    add_card(s1, Inches(1.0), Inches(4.5), Inches(3.6), Inches(2.2), "Project Domain", [
        "Computer Vision & Deep Learning",
        "Kinematic Joint Biomechanics",
        "Sports Healthcare Analytics"
    ], C_BLUE)
    
    add_card(s1, Inches(4.85), Inches(4.5), Inches(3.6), Inches(2.2), "System Architecture", [
        "FastAPI & OpenCV Python Backend",
        "MediaPipe 33-Landmark Pose ML",
        "React 18 & Vite Interactive HUD"
    ], C_GREEN)
    
    add_card(s1, Inches(8.7), Inches(4.5), Inches(3.6), Inches(2.2), "Project Status", [
        "Completed All 4 Milestones",
        "100% RBAC Security Verified",
        "Production-Ready & Fully Deployed"
    ], C_AMBER)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Motivation
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Statement", "The Challenge of Non-Contact Athletic Injuries", 2)
    
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.7), Inches(3.8), "High Non-Contact Injury Rates", [
        "Over 70% of ACL tears and hamstring strains occur during non-contact deceleration, landing, and rapid cutting motions.",
        "Significant medical, surgical, and rehabilitation financial losses across collegiate and professional sports organizations.",
        "Subtle joint misalignments (e.g. knee valgus collapse) occur in under 50ms, evading human visual coaching observation."
    ], C_RED)
    
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(3.8), "Limitations of Existing Methods", [
        "Manual Visual Coaching: Subjective, qualitative, and impossible to measure micro-deviations during live motion.",
        "Laboratory Motion Capture: Costs exceed $50,000+, requires intrusive retro-reflective physical skin markers, and cannot be used outdoors.",
        "Disconnected Workload Records: Training load fatigue and kinematic movement telemetry are rarely unified."
    ], C_AMBER)
    
    # Callout
    c_box = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.65), Inches(11.7), Inches(1.2))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = C_CARD
    c_box.line.color.rgb = C_BLUE
    c_box.line.width = Pt(1.5)
    ctf = c_box.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.25)
    ctf.margin_top = Inches(0.18)
    cp = ctf.paragraphs[0]
    cp.text = "CORE OBJECTIVE: Deliver an automated, markerless, video-based assessment platform that extracts 3D skeleton keypoints, calculates joint angles, computes acute fatigue ratios, and generates clinical evaluation reports in real time."
    cp.font.name = 'Arial'
    cp.font.size = Pt(11)
    cp.font.bold = True
    cp.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 3: Objectives & Scope
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "Project Objectives", "Key Project Engineering Objectives", 3)
    
    w_4 = Inches(2.78)
    add_card(s3, Inches(0.8), Inches(1.6), w_4, Inches(5.2), "1. Biometrics & RBAC", [
        "Multi-role user authentication (Athlete, Coach, Admin, Physio).",
        "Comprehensive athlete profiles (biometrics, sports, medical history, disability flags).",
        "Strict multi-tenant coach-athlete isolation.",
        "User profile self-service editing."
    ], C_BLUE)
    
    add_card(s3, Inches(3.78), Inches(1.6), w_4, Inches(5.2), "2. Pose & Tracking", [
        "Markerless 33-keypoint 3D anatomical landmark detection.",
        "MediaPipe Pose & OpenCV frame pipeline.",
        "Real-time skeleton wireframe overlay rendering.",
        "Interactive joint angle telemetry player."
    ], C_GREEN)
    
    add_card(s3, Inches(6.76), Inches(1.6), w_4, Inches(5.2), "3. Biomechanical AI", [
        "Continuous 3D joint angle vector calculations (Valgus, Flexion, Spine).",
        "Bilateral limb asymmetry detection.",
        "Acute-to-Chronic Workload Ratio (ACWR) fatigue modeling.",
        "Composite Injury Risk Index (0–100%)."
    ], C_AMBER)
    
    add_card(s3, Inches(9.74), Inches(1.6), w_4, Inches(5.2), "4. Reports & Logs", [
        "Executive Clinical Evaluation Reports with printable summaries.",
        "Algorithmic personalized corrective rehabilitation drills.",
        "Date-range filtered historical session directory.",
        "Proactive high-risk threshold alert notifications."
    ], C_RED)

    # -------------------------------------------------------------
    # SLIDE 4: System Architecture
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "System Architecture", "End-to-End Integrated 4-Milestone Pipeline", 4)
    
    add_card(s4, Inches(0.8), Inches(1.6), w_4, Inches(4.3), "Milestone 1\nAuth & Biometrics", [
        "FastAPI JWT authentication.",
        "Athlete registry in MongoDB.",
        "Dynamic Coach selection.",
        "Role-based permission gating."
    ], C_BLUE)
    
    add_card(s4, Inches(3.78), Inches(1.6), w_4, Inches(4.3), "Milestone 2\nVideo & Skeleton", [
        "Multi-format video ingestion.",
        "OpenCV frame decoding.",
        "MediaPipe 33-landmark model.",
        "Real-time kinematic wireframe."
    ], C_GREEN)
    
    add_card(s4, Inches(6.76), Inches(1.6), w_4, Inches(4.3), "Milestone 3\nAthlete Intelligence", [
        "Kinematic anomaly engine.",
        "Limb asymmetry matrix.",
        "ACWR fatigue calculation.",
        "Biomechanical index scoring."
    ], C_AMBER)
    
    add_card(s4, Inches(9.74), Inches(1.6), w_4, Inches(4.3), "Milestone 4\nReports & History", [
        "Clinical evaluation summary.",
        "Corrective drill generation.",
        "Date-filtered session logs.",
        "Real-time alert broadcasting."
    ], C_RED)
    
    # Unified Host Callout
    c_box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = C_CARD
    c_box.line.color.rgb = C_GREEN
    ctf = c_box.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.2)
    ctf.margin_top = Inches(0.12)
    cp = ctf.paragraphs[0]
    cp.text = "UNIFIED APPLICATION HOST: All 4 milestones execute on a single React 18 frontend (http://localhost:5173/) backed by a high-throughput async FastAPI service (http://127.0.0.1:8000/)."
    cp.font.name = 'Arial'
    cp.font.size = Pt(11)
    cp.font.bold = True
    cp.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 5: Methodology & Mathematical Model
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "Methodology", "Mathematical & Biomechanical Modeling Engine", 5)
    
    w_3 = Inches(3.75)
    add_card(s5, Inches(0.8), Inches(1.6), w_3, Inches(5.2), "1. 3D Joint Angle Calculation", [
        "Computes 3D spatial vector dot product across adjacent anatomical landmark triples (e.g. Hip-Knee-Ankle):",
        "Formula:  θ = arccos( (u · v) / (||u|| ||v||) )",
        "Measures continuous Knee Valgus, Hip Flexion, Ankle Dorsiflexion, and Spine Lateral Sway.",
        "Identifies excessive inward knee collapse angles (>15°) linked to ACL tears."
    ], C_BLUE)
    
    add_card(s5, Inches(4.78), Inches(1.6), w_3, Inches(5.2), "2. Acute:Chronic Workload (ACWR)", [
        "Models short-term athletic fatigue versus long-term physical conditioning:",
        "Formula:  ACWR = Acute Load (7-Day) / Chronic Load (28-Day)",
        "Optimal Sweet Spot: 0.80 – 1.30 (low injury risk, progressive fitness).",
        "Danger Zone: > 1.50 (fatigue spike, soft-tissue injury risk increases by 3-5x)."
    ], C_GREEN)
    
    add_card(s5, Inches(8.75), Inches(1.6), w_3, Inches(5.2), "3. Composite Risk Index", [
        "Synthesizes kinematics, asymmetry, ACWR, and medical history into a weighted 0–100% score:",
        "Formula:  Risk = w1·Dev + w2·Asym + w3·ACWR + w4·History",
        "Low Risk: 0% – 35% (Cleared)",
        "Moderate Risk: 36% – 70% (Corrective drills)",
        "Critical Risk: > 70% (Intervention required)"
    ], C_RED)

    # -------------------------------------------------------------
    # SLIDE 6: Technologies Used
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "Technology Stack", "Full-Stack & Computer Vision Technology Ecosystem", 6)
    
    add_card(s6, Inches(0.8), Inches(1.6), w_3, Inches(5.2), "Frontend Technologies", [
        "React 18 & Vite: High-performance single-page app architecture with Hot Module Reloading.",
        "Tailwind CSS v4: Custom HUD dark & high-contrast light theme token system.",
        "Lucide React: Crisp vector iconography.",
        "Framer Motion: Hardware-accelerated transitions.",
        "Axios & Context API: Global session JWT state."
    ], C_BLUE)
    
    add_card(s6, Inches(4.78), Inches(1.6), w_3, Inches(5.2), "Backend & Vision Pipeline", [
        "FastAPI & Python 3.10: High-concurrency async REST API framework.",
        "OpenCV (cv2): Video file validation, frame rate normalization, and stream processing.",
        "MediaPipe Pose: Google's 33-landmark 3D spatial pose estimation model.",
        "NumPy & SciPy: Vectorized matrix mathematics for joint angle and asymmetry computation."
    ], C_GREEN)
    
    add_card(s6, Inches(8.75), Inches(1.6), w_3, Inches(5.2), "Database & Security", [
        "MongoDB & Motor: High-throughput non-blocking document database.",
        "JWT & Bcrypt: Secure salted password encryption and cryptographic session tokens.",
        "Role-Based Access Control: Strict data boundary enforcement between coaches and athletes.",
        "Uvicorn: Lightning-fast ASGI event-loop server."
    ], C_AMBER)

    # -------------------------------------------------------------
    # SLIDE 7: Milestone 1 & 2 Implementation
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "Implementation", "Milestone 1 & 2: Ingestion, Biometrics & Tracking", 7)
    
    w_2 = Inches(5.7)
    add_card(s7, Inches(0.8), Inches(1.6), w_2, Inches(5.2), "Milestone 1: Authentication & Athlete Hub", [
        "Role-Based Authentication: Granular login/registration for Athlete, Coach, Admin, and Physiotherapist roles.",
        "Biometric Athlete Profiles: Tracks Age, Gender, Height, Weight, Sport, Playing Position, Load, Experience, Medical Notes, and Disability Status.",
        "Coach Selection & Association: Dynamic coach assignment stored in MongoDB.",
        "Self-Service Profile Updates: All authenticated users can edit their name, email, and password via PUT /me."
    ], C_BLUE)
    
    add_card(s7, Inches(6.8), Inches(1.6), w_2, Inches(5.2), "Milestone 2: Video Ingestion & Skeleton Analytics", [
        "Video File Ingestion: MP4/MOV validation with video deletion and server disk cleanup.",
        "MediaPipe 33-Keypoint Extraction: High-precision anatomical landmark tracking per video frame.",
        "Real-Time Skeleton Tracking: Renders dynamic 2D/3D kinetic skeleton wireframe overlays.",
        "Interactive Kinematics Player: Frame-by-frame joint angle telemetry playback."
    ], C_GREEN)

    # -------------------------------------------------------------
    # SLIDE 8: Milestone 3 & 4 Implementation
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8)
    add_header(s8, "Implementation", "Milestone 3 & 4: AI Intelligence & Reporting", 8)
    
    add_card(s8, Inches(0.8), Inches(1.6), w_2, Inches(5.2), "Milestone 3: Athlete Intelligence Engine", [
        "Biomechanical Vulnerability Analysis: Identifies high-risk joint stress patterns (knee valgus collapse, trunk sway).",
        "Bilateral Asymmetry Matrix: Evaluates left vs right limb mechanical imbalances.",
        "ACWR Fatigue Monitoring: Tracks cumulative workload ratios to predict overtraining strain.",
        "Dynamic Anomaly Badging: Automatically flags joints requiring clinical attention."
    ], C_AMBER)
    
    add_card(s8, Inches(6.8), Inches(1.6), w_2, Inches(5.2), "Milestone 4: Evaluation Reports & History", [
        "Executive Athlete Evaluation Report: Clinical diagnostic summary with risk classification badges and observations.",
        "Targeted Corrective Drills: Algorithmic prescription of customized rehabilitation exercises (Banded Squats, Nordic Curls).",
        "Date-Range Filtered Analysis Logs: Synchronized session directory querying records between Start and End dates.",
        "Proactive Anomaly Alerts: Real-time notification dispatching for high-risk movement sessions."
    ], C_RED)

    # -------------------------------------------------------------
    # SLIDE 9: Key Features & UI Capabilities
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9)
    add_header(s9, "Key Features", "User Experience & Platform Capabilities", 9)
    
    add_card(s9, Inches(0.8), Inches(1.6), w_3, Inches(2.5), "1. Unified Single-Port Flow", [
        "Login → Dashboard → Video Upload → Skeleton Tracking → Intelligence → Report → History with zero re-logins."
    ], C_BLUE)
    
    add_card(s9, Inches(4.78), Inches(1.6), w_3, Inches(2.5), "2. Dual-Theme Accessibility", [
        "Full HUD Dark mode and clean Light mode with high-contrast slate text and no invisible elements."
    ], C_GREEN)
    
    add_card(s9, Inches(8.75), Inches(1.6), w_3, Inches(2.5), "3. Multi-Tenant Isolation", [
        "Coaches only see assigned athletes. Cross-coach access by URL or API is blocked with 403 Forbidden."
    ], C_AMBER)
    
    add_card(s9, Inches(0.8), Inches(4.3), w_3, Inches(2.5), "4. Video Lifecycle Controls", [
        "Upload, preview, analyze, and safely delete video sessions with permanent server file removal."
    ], C_RED)
    
    add_card(s9, Inches(4.78), Inches(4.3), w_3, Inches(2.5), "5. Date-Range Historical Logs", [
        "Filter analysis sessions by exact start and end dates with validation and one-click report review."
    ], C_BLUE)
    
    add_card(s9, Inches(8.75), Inches(4.3), w_3, Inches(2.5), "6. Self-Service Profile Updates", [
        "All authenticated roles (Athlete, Coach, Admin) can update their name, email, and password directly."
    ], C_GREEN)

    # -------------------------------------------------------------
    # SLIDE 10: Performance, Testing & Verification
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s10)
    add_header(s10, "Testing & Verification", "System Validation, Security & Performance Metrics", 10)
    
    add_card(s10, Inches(0.8), Inches(1.6), w_2, Inches(5.2), "Functional & Security Test Results", [
        "Coach-Athlete Data Isolation: 100% pass on Coach A vs Coach B cross-access isolation tests.",
        "Authorization Gating: Milestone 3 & 4 APIs reject unauthorized session queries with 403 Forbidden.",
        "User Profile API: Verified successful name, email, and password updates via PUT /me.",
        "Video Lifecycle: Verified video upload, MediaPipe pose extraction, and disk file erasure upon deletion.",
        "Date Filter Validation: Verified start_date <= end_date query constraints."
    ], C_GREEN)
    
    add_card(s10, Inches(6.8), Inches(1.6), w_2, Inches(5.2), "Performance & Build Benchmarks", [
        "Vite Production Bundle: ✓ built in 2.04s (6,278 modules transformed, 0 errors, 0 warnings).",
        "Pose Inference Throughput: ~28–30 FPS real-time tracking throughput on multi-core CPU.",
        "FastAPI Latency: Sub-15ms average response time on session and history endpoints.",
        "Database Persistence: 100% real MongoDB persistence across all collections with zero mock data."
    ], C_BLUE)

    # -------------------------------------------------------------
    # SLIDE 11: Challenges & Solutions
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s11)
    add_header(s11, "Challenges & Solutions", "Key Engineering Challenges & Technical Resolutions", 11)
    
    add_card(s11, Inches(0.8), Inches(1.6), w_2, Inches(2.5), "1. Multi-Host Fragmentation", [
        "Problem: Milestones 3 & 4 were on isolated ports (5174/5175).",
        "Solution: Consolidated into a unified single-port React app on port 5173 with shared JWT context."
    ], C_AMBER)
    
    add_card(s11, Inches(6.8), Inches(1.6), w_2, Inches(2.5), "2. Multi-Tenant Coach Isolation", [
        "Problem: Preventing coaches from seeing another coach's athletes.",
        "Solution: Implemented backend ownership validation across all routers returning 403 Forbidden."
    ], C_BLUE)
    
    add_card(s11, Inches(0.8), Inches(4.3), w_2, Inches(2.5), "3. Light-Theme Text Contrast", [
        "Problem: Inverted backgrounds caused white-on-white text illegibility.",
        "Solution: Built a global CSS variable token system ensuring dark, high-contrast text across all cards."
    ], C_GREEN)
    
    add_card(s11, Inches(6.8), Inches(4.3), w_2, Inches(2.5), "4. File Lifecycle & Storage", [
        "Problem: Deleting video records left orphaned files on disk.",
        "Solution: Added filesystem unlinking and cascading MongoDB session cleanup on delete."
    ], C_RED)

    # -------------------------------------------------------------
    # SLIDE 12: Future Scope & Conclusion
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s12)
    add_header(s12, "Future Scope & Conclusion", "Roadmap & Project Summary", 12)
    
    add_card(s12, Inches(0.8), Inches(1.6), w_2, Inches(5.2), "Future Roadmap", [
        "Wearable Sensor Fusion: Integrate IMU sensor streams (accelerometers, gyroscopes) with vision coordinates.",
        "Live WebRTC Stream Analysis: Real-time on-pitch sideline injury risk alerts during competitive matches.",
        "Multi-Athlete Tracking: Simultaneous multi-player skeletal tracking for team sports.",
        "Deep Sequential Forecasting: Transformer-based temporal models for long-term injury forecasting."
    ], C_BLUE)
    
    add_card(s12, Inches(6.8), Inches(1.6), w_2, Inches(5.2), "Conclusion", [
        "Successfully developed and verified KineticGuard, a production-grade AI sports injury platform.",
        "Democratized clinical-grade biomechanical analytics without expensive $50,000+ motion labs.",
        "Accomplished end-to-end integration across all 4 milestones with robust RBAC and high UX fidelity.",
        "Verified all test suites with zero errors and production-ready architecture."
    ], C_GREEN)

    # -------------------------------------------------------------
    # SLIDE 13: Thank You & Q&A
    # -------------------------------------------------------------
    s13 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s13)
    
    ty_box = s13.shapes.add_textbox(Inches(2.0), Inches(2.2), Inches(9.333), Inches(3.0))
    tyf = ty_box.text_frame
    tyf.word_wrap = True
    
    p_ty1 = tyf.paragraphs[0]
    p_ty1.alignment = PP_ALIGN.CENTER
    p_ty1.text = "Thank You!"
    p_ty1.font.name = 'Arial'
    p_ty1.font.size = Pt(44)
    p_ty1.font.bold = True
    p_ty1.font.color.rgb = C_BLUE
    p_ty1.space_after = Pt(14)
    
    p_ty2 = tyf.add_paragraph()
    p_ty2.alignment = PP_ALIGN.CENTER
    p_ty2.text = "KineticGuard: AI-Powered Sports Injury Risk Assessment Platform"
    p_ty2.font.name = 'Arial'
    p_ty2.font.size = Pt(18)
    p_ty2.font.bold = True
    p_ty2.font.color.rgb = C_WHITE
    p_ty2.space_after = Pt(10)
    
    p_ty3 = tyf.add_paragraph()
    p_ty3.alignment = PP_ALIGN.CENTER
    p_ty3.text = "Open for Questions & Discussion"
    p_ty3.font.name = 'Arial'
    p_ty3.font.size = Pt(15)
    p_ty3.font.color.rgb = C_MUTED
    
    # Save Presentation
    out_dir_1 = r"c:\Users\Nagaveni\OneDrive\Desktop\Sports injury"
    out_dir_2 = r"c:\Users\Nagaveni\OneDrive\Desktop\Sports injury\sports-injury-risk"
    
    fpath_1 = os.path.join(out_dir_1, "KineticGuard_Project_Presentation.pptx")
    fpath_2 = os.path.join(out_dir_2, "KineticGuard_Project_Presentation.pptx")
    
    prs.save(fpath_1)
    prs.save(fpath_2)
    print(f"Presentation saved successfully to:\n1. {fpath_1}\n2. {fpath_2}")

if __name__ == "__main__":
    create_presentation()

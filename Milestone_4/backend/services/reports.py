"""
Milestone 4 — Reports & Export System (PDF section 12)
Location: backend/services/reports.py

Generates the report types the PDF lists:
  - Injury risk reports
  - Biomechanical assessment reports
  - Movement analysis reports
  - Athlete performance reports
(Rehabilitation reports use the same recommendations data — there's no
separate rehab-plan-tracking model yet, so "rehabilitation report" and
"injury risk report" render from the same underlying data for now.)

Two export formats, both from the PDF ("PDF export", "Excel export"):
  - generate_video_pdf_report()   -> single video: biomechanics + risk assessment
  - generate_athlete_excel_report() -> athlete performance workbook across all videos
  - generate_athlete_pdf_report()  -> athlete summary PDF across all videos
"""

import io
from datetime import datetime
from typing import Any, List, Optional

from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BRAND_COLOR = colors.HexColor("#8b5cf6")
RISK_COLORS = {
    "Low": colors.HexColor("#16a34a"),
    "Moderate": colors.HexColor("#d97706"),
    "High": colors.HexColor("#e11d48"),
    "Critical": colors.HexColor("#991b1b"),
}


def _styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle", parent=styles["Title"], textColor=BRAND_COLOR, spaceAfter=6
        )
    )
    styles.add(
        ParagraphStyle(
            name="SectionHeading", parent=styles["Heading2"], textColor=BRAND_COLOR, spaceBefore=14, spaceAfter=6
        )
    )
    styles.add(ParagraphStyle(name="Small", parent=styles["Normal"], fontSize=9, textColor=colors.grey))
    return styles


def _kv_table(rows: List[List[str]]) -> Table:
    table = Table(rows, colWidths=[2.4 * inch, 3.6 * inch])
    table.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#6b7280")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#f3e8ff")),
            ]
        )
    )
    return table


def _header_block(styles, title: str, athlete_owner_name: str, sport: str) -> list:
    flow = [
        Paragraph(title, styles["ReportTitle"]),
        Paragraph(f"{athlete_owner_name} &middot; {sport}", styles["Normal"]),
        Paragraph(
            f"Generated {datetime.utcnow().strftime('%B %d, %Y %H:%M UTC')}", styles["Small"]
        ),
        Spacer(1, 14),
    ]
    return flow


def generate_video_pdf_report(
    video: Any,
    athlete: Any,
    owner_name: str,
    biomechanics_report: Optional[dict],
    risk_assessment: Optional[dict],
) -> io.BytesIO:
    """Injury risk report + Biomechanical assessment report + Movement
    analysis report for a single video (PDF section 12)."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, topMargin=0.8 * inch, bottomMargin=0.8 * inch)
    styles = _styles()
    flow = _header_block(styles, "Movement Analysis Report", owner_name, athlete.sport if athlete else "")

    flow.append(Paragraph("Video Details", styles["SectionHeading"]))
    flow.append(
        _kv_table(
            [
                ["Activity type", (video.activity_type or "-").replace("_", " ").title()],
                ["Uploaded", video.uploaded_at.strftime("%B %d, %Y") if video.uploaded_at else "-"],
                ["Duration", f"{video.duration_seconds:.1f}s" if video.duration_seconds else "-"],
                ["Status", video.status],
            ]
        )
    )

    if biomechanics_report:
        flow.append(Paragraph("Biomechanical Assessment", styles["SectionHeading"]))
        flow.append(
            _kv_table(
                [
                    ["Movement quality score", str(biomechanics_report.get("movement_quality_score", "-"))],
                    ["Avg trunk lean", f"{biomechanics_report.get('avg_trunk_lean', '-')}\u00b0"],
                    ["Knee valgus asymmetry", str(biomechanics_report.get("knee_valgus_asymmetry", "-"))],
                    ["Movement symmetry score", str(biomechanics_report.get("movement_symmetry_score", "-"))],
                ]
            )
        )
    else:
        flow.append(Paragraph("Biomechanical Assessment", styles["SectionHeading"]))
        flow.append(Paragraph("Not available for this video.", styles["Normal"]))

    if risk_assessment:
        flow.append(Paragraph("Injury Risk Prediction", styles["SectionHeading"]))
        category = risk_assessment.get("risk_category", "-")
        color = RISK_COLORS.get(category, colors.black)
        risk_style = ParagraphStyle(
            name="RiskValue", parent=styles["Normal"], textColor=color, fontSize=14, fontName="Helvetica-Bold"
        )
        flow.append(
            Paragraph(
                f"{category} Risk &mdash; {risk_assessment.get('overall_injury_risk_score', '-')}/100",
                risk_style,
            )
        )
        flow.append(Spacer(1, 6))
        flow.append(
            _kv_table(
                [
                    ["Biomechanical deviations (35%)", str(risk_assessment.get("biomechanical_deviation_score", "-"))],
                    ["Historical injury factors (20%)", str(risk_assessment.get("historical_injury_score", "-"))],
                    ["Movement asymmetry (20%)", str(risk_assessment.get("movement_asymmetry_score", "-"))],
                    ["Training load indicators (15%)", str(risk_assessment.get("training_load_score", "-"))],
                    ["Fatigue indicators (10%)", str(risk_assessment.get("fatigue_score", "-"))],
                    ["Overall athlete health score", str(risk_assessment.get("overall_athlete_health_score", "-"))],
                ]
            )
        )

        injury_type_risks = risk_assessment.get("injury_type_risks") or {}
        if injury_type_risks:
            flow.append(Paragraph("Risk by Injury Category", styles["SectionHeading"]))
            rows = [["Category", "Risk /100"]] + [[k, str(v)] for k, v in injury_type_risks.items()]
            cat_table = Table(rows, colWidths=[3.2 * inch, 1.8 * inch])
            cat_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ede9fe")),
                        ("FONTSIZE", (0, 0), (-1, -1), 10),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#f3e8ff")),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            flow.append(cat_table)

        anomalies = risk_assessment.get("anomalies_detected") or []
        if anomalies:
            flow.append(Paragraph("Movement Anomalies Detected", styles["SectionHeading"]))
            for a in anomalies:
                flow.append(Paragraph(f"&bull; {a.get('description', '')}", styles["Normal"]))

        recommendations = risk_assessment.get("recommendations") or []
        if recommendations:
            flow.append(Paragraph("Corrective Recommendations", styles["SectionHeading"]))
            for rec in recommendations:
                flow.append(
                    Paragraph(f"<b>{rec.get('title', '')}</b> &mdash; {rec.get('description', '')}", styles["Normal"])
                )
                flow.append(Spacer(1, 4))
    else:
        flow.append(Paragraph("Injury Risk Prediction", styles["SectionHeading"]))
        flow.append(Paragraph("Not available for this video.", styles["Normal"]))

    flow.append(Spacer(1, 20))
    flow.append(
        Paragraph(
            "This is a heuristic screening report, not a medical diagnosis. "
            "Consult a physiotherapist or physician for clinical decisions.",
            styles["Small"],
        )
    )

    doc.build(flow)
    buffer.seek(0)
    return buffer


def generate_athlete_pdf_report(
    athlete: Any,
    owner_name: str,
    video_rows: List[dict],
) -> io.BytesIO:
    """Athlete performance report (PDF section 12) — summary across every
    analyzed video for this athlete."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, topMargin=0.8 * inch, bottomMargin=0.8 * inch)
    styles = _styles()
    flow = _header_block(styles, "Athlete Performance Report", owner_name, athlete.sport if athlete else "")

    flow.append(Paragraph("Athlete Profile", styles["SectionHeading"]))
    flow.append(
        _kv_table(
            [
                ["Sport", athlete.sport or "-"],
                ["Position", athlete.position or "-"],
                ["Age", str(athlete.age) if athlete.age is not None else "-"],
                ["Height / Weight", f"{athlete.height or '-'} cm / {athlete.weight or '-'} kg"],
                ["Training load", athlete.training_load or "-"],
            ]
        )
    )

    if video_rows:
        scores = [r["overall_injury_risk_score"] for r in video_rows if r.get("overall_injury_risk_score") is not None]
        avg_score = round(sum(scores) / len(scores), 1) if scores else None
        flow.append(Paragraph("Summary", styles["SectionHeading"]))
        flow.append(
            _kv_table(
                [
                    ["Videos analyzed", str(len(video_rows))],
                    ["Average injury risk score", str(avg_score) if avg_score is not None else "-"],
                ]
            )
        )

        flow.append(Paragraph("Risk History", styles["SectionHeading"]))
        rows = [["Date", "Activity", "Risk Score", "Category"]]
        for r in video_rows:
            rows.append(
                [
                    r["uploaded_at"].strftime("%b %d, %Y") if r.get("uploaded_at") else "-",
                    (r.get("activity_type") or "-").replace("_", " ").title(),
                    str(r.get("overall_injury_risk_score", "-")),
                    r.get("risk_category", "-"),
                ]
            )
        history_table = Table(rows, colWidths=[1.3 * inch, 1.8 * inch, 1.1 * inch, 1.1 * inch])
        history_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ede9fe")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#f3e8ff")),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        flow.append(history_table)
    else:
        flow.append(Paragraph("No analyzed videos yet for this athlete.", styles["Normal"]))

    doc.build(flow)
    buffer.seek(0)
    return buffer


def generate_athlete_excel_report(
    athlete: Any,
    owner_name: str,
    video_rows: List[dict],
) -> io.BytesIO:
    """Athlete performance report as an .xlsx workbook (PDF "Excel export")."""
    wb = Workbook()

    header_fill = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")
    header_font = Font(bold=True, color="4C1D95")
    title_font = Font(bold=True, size=14, color="6D28D9")

    # --- Sheet 1: Overview ---
    ws1 = wb.active
    ws1.title = "Overview"
    ws1["A1"] = "Athlete Performance Report"
    ws1["A1"].font = title_font
    ws1["A2"] = f"{owner_name} · {athlete.sport if athlete else ''}"
    ws1["A3"] = f"Generated {datetime.utcnow().strftime('%B %d, %Y %H:%M UTC')}"

    overview_fields = [
        ("Sport", athlete.sport if athlete else "-"),
        ("Position", athlete.position if athlete else "-"),
        ("Age", athlete.age if athlete else "-"),
        ("Height (cm)", athlete.height if athlete else "-"),
        ("Weight (kg)", athlete.weight if athlete else "-"),
        ("Training load", athlete.training_load if athlete else "-"),
        ("Videos analyzed", len(video_rows)),
    ]
    row = 5
    for label, value in overview_fields:
        ws1.cell(row=row, column=1, value=label).font = Font(bold=True, color="6B7280")
        ws1.cell(row=row, column=2, value=value)
        row += 1
    ws1.column_dimensions["A"].width = 24
    ws1.column_dimensions["B"].width = 30

    # --- Sheet 2: Risk History ---
    ws2 = wb.create_sheet("Risk History")
    headers = [
        "Date",
        "Activity",
        "Overall Risk Score",
        "Risk Category",
        "Biomechanical Deviation",
        "Historical Injury",
        "Movement Asymmetry",
        "Training Load",
        "Fatigue",
    ]
    for col, h in enumerate(headers, start=1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    for i, r in enumerate(video_rows, start=2):
        ws2.cell(row=i, column=1, value=r["uploaded_at"].strftime("%Y-%m-%d") if r.get("uploaded_at") else "-")
        ws2.cell(row=i, column=2, value=(r.get("activity_type") or "-").replace("_", " ").title())
        ws2.cell(row=i, column=3, value=r.get("overall_injury_risk_score"))
        ws2.cell(row=i, column=4, value=r.get("risk_category"))
        ws2.cell(row=i, column=5, value=r.get("biomechanical_deviation_score"))
        ws2.cell(row=i, column=6, value=r.get("historical_injury_score"))
        ws2.cell(row=i, column=7, value=r.get("movement_asymmetry_score"))
        ws2.cell(row=i, column=8, value=r.get("training_load_score"))
        ws2.cell(row=i, column=9, value=r.get("fatigue_score"))
    for col_letter in "ABCDEFGHI":
        ws2.column_dimensions[col_letter].width = 18

    # --- Sheet 3: Injury Category Breakdown (latest assessment only) ---
    ws3 = wb.create_sheet("Injury Categories")
    ws3.cell(row=1, column=1, value="Injury Category").font = header_font
    ws3.cell(row=1, column=2, value="Risk /100 (latest video)").font = header_font
    ws3.cell(row=1, column=1).fill = header_fill
    ws3.cell(row=1, column=2).fill = header_fill
    latest_types = video_rows[0].get("injury_type_risks") if video_rows else None
    if latest_types:
        for i, (k, v) in enumerate(latest_types.items(), start=2):
            ws3.cell(row=i, column=1, value=k)
            ws3.cell(row=i, column=2, value=v)
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 22

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
